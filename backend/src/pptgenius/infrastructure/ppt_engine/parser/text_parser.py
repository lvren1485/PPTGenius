"""Text element renderer."""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from .base import TextboxElement
from .styles import (
    set_text_strikethrough,
    set_text_small_caps,
    set_text_all_caps,
    set_text_kerning,
    set_text_shadow,
    set_text_glow,
    apply_shape_fill,
    apply_shape_line,
)


def render_textbox(slide, el: TextboxElement) -> None:
    """Render a textbox element onto a slide."""
    left = Inches(el.position.left)
    top = Inches(el.position.top)
    width = Inches(el.position.width)
    height = Inches(el.position.height) if el.position.height else Inches(1.0)

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP

    for i, block in enumerate(el.content):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _apply_paragraph(p, block.paragraph)

    # Optional rotation
    if el.rotation:
        txBox.rotation = el.rotation

    # Optional fill
    if el.fill:
        apply_shape_fill(txBox, el.fill.type, el.fill.color)

    # Optional line
    if el.line and el.line.color:
        apply_shape_line(txBox, el.line.color, el.line.width_pt, el.line.dash)


ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}


def _apply_paragraph(p, para_spec) -> None:
    """Apply paragraph-level formatting."""
    p.alignment = ALIGN_MAP.get(para_spec.alignment, PP_ALIGN.LEFT)
    p.level = para_spec.level

    if para_spec.space_before_pt is not None:
        p.space_before = Pt(para_spec.space_before_pt)
    if para_spec.space_after_pt is not None:
        p.space_after = Pt(para_spec.space_after_pt)

    for run_spec in para_spec.runs:
        run = p.add_run()
        run.text = run_spec.text
        if run_spec.font:
            _apply_font(run, run_spec.font)


def _apply_font(run, font_spec) -> None:
    """Apply run-level font formatting."""
    f = run.font
    if font_spec.name:
        f.name = font_spec.name
    if font_spec.size:
        f.size = Pt(font_spec.size)
    if font_spec.bold is not None:
        f.bold = font_spec.bold
    if font_spec.italic is not None:
        f.italic = font_spec.italic
    if font_spec.color:
        f.color.rgb = RGBColor.from_string(font_spec.color)
    if font_spec.underline:
        from pptx.enum.text import MSO_TEXT_UNDERLINE_TYPE
        ul_map = {
            "none": MSO_TEXT_UNDERLINE_TYPE.NONE,
            "single": MSO_TEXT_UNDERLINE_TYPE.SINGLE_LINE,
            "double": MSO_TEXT_UNDERLINE_TYPE.DOUBLE_LINE,
            "wavy": MSO_TEXT_UNDERLINE_TYPE.WAVY_LINE,
        }
        f.underline = ul_map.get(font_spec.underline, MSO_TEXT_UNDERLINE_TYPE.SINGLE_LINE)

    # lxml-based effects
    if font_spec.strikethrough:
        set_text_strikethrough(run, "sngStrike" if font_spec.strikethrough == "single" else "dblStrike")
    if font_spec.small_caps:
        set_text_small_caps(run)
    if font_spec.all_caps:
        set_text_all_caps(run)
    if font_spec.kerning_pt is not None:
        set_text_kerning(run, font_spec.size or 12, font_spec.kerning_pt)
    if font_spec.shadow:
        set_text_shadow(
            run,
            blur_pt=font_spec.shadow.blur_pt,
            offset_pt=font_spec.shadow.offset_pt,
            angle_deg=font_spec.shadow.angle_deg,
            color_hex=font_spec.shadow.color,
            alpha_pct=font_spec.shadow.alpha_pct,
        )
    if font_spec.glow:
        set_text_glow(
            run,
            radius_pt=font_spec.glow.radius_pt,
            color_hex=font_spec.glow.color,
            alpha_pct=font_spec.glow.alpha_pct,
        )
